"""
Beginner Level Task — Retail Store Sales Analysis
ShadowFox Data Analyst Internship

What this script does (in order):
  1. Generates the raw dataset (with realistic mess left in, on purpose)
  2. Cleans it and saves the cleaned version
  3. Prints the cleaning audit + key metrics
  4. Builds the Excel dashboard (KPI cards + charts)
  5. Renders a PNG preview of the dashboard layout

Run:  python beginner_task.py
Needs: pandas, numpy, matplotlib, openpyxl
"""

import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

SEED = 42
rng = np.random.default_rng(SEED)

BASE = os.path.dirname(os.path.abspath(__file__))
RAW_DIR  = os.path.join(BASE, "data", "raw")
CLEAN_DIR = os.path.join(BASE, "data", "cleaned")
DASH_DIR = os.path.join(BASE, "dashboard")
for d in (RAW_DIR, CLEAN_DIR, DASH_DIR):
    os.makedirs(d, exist_ok=True)

# ----------------------------------------------------------------------
# 1. GENERATE THE RAW DATASET  (Jan-Dec 2025, 4 categories, 4 regions)
# ----------------------------------------------------------------------
months = pd.date_range("2025-01-01", "2025-12-01", freq="MS")

# (mean price, price spread) per category
cats = {"Electronics": (120, 45), "Clothing": (35, 12),
        "Groceries": (12, 4), "Home & Living": (55, 20)}
regions = ["North", "South", "East", "West"]

rows = []
order_id = 1000
for m in months:
    # seasonality: gentle wave + a festive lift in Oct/Nov/Dec
    seasonal = 1 + 0.35 * np.sin(2 * np.pi * (m.month - 3) / 12) \
                  + (0.5 if m.month in [10, 11, 12] else 0)
    for _ in range(int(45 * seasonal)):
        cat = rng.choice(list(cats), p=[.25, .30, .30, .15])
        mean, sd = cats[cat]
        price = round(max(3, rng.normal(mean, sd)), 2)
        qty = int(rng.integers(1, 5))
        rows.append([f"ORD-{order_id}", m.strftime("%Y-%m-%d"), cat,
                     rng.choice(regions), price, qty, round(price * qty, 2)])
        order_id += 1

df = pd.DataFrame(rows, columns=["Order_ID", "Order_Date", "Category",
                                 "Region", "Unit_Price", "Quantity", "Total_Sales"])

# Plant realistic mess (this is the "raw export"):
raw = df.copy()
raw.loc[raw.sample(8, random_state=1).index, "Region"] = None                      # blank regions
raw.loc[raw.sample(5, random_state=2).index, "Unit_Price"] *= -1                   # refund-entry errors
raw.loc[raw.sample(4, random_state=3).index, "Category"] = \
    raw["Category"].str.upper()                                                    # ELECTRONICS etc.

raw_path = os.path.join(RAW_DIR, "retail_store_sales_raw.csv")
raw.to_csv(raw_path, index=False)

# ----------------------------------------------------------------------
# 2. CLEANING  (audit first, then fix)
# ----------------------------------------------------------------------
print("== CLEANING AUDIT (raw file) ==")
print("rows:", len(raw),
      "| blank Region:", raw["Region"].isna().sum(),
      "| negative Unit_Price:", (raw["Unit_Price"] < 0).sum(),
      "| bad Category casing:", (raw["Category"] != raw["Category"].str.title()).sum())

clean = raw.copy()

# 2a. blank Region -> that month's most common region (reasonable guess)
clean["Region"] = clean.groupby("Order_Date")["Region"] \
                       .transform(lambda s: s.fillna(s.mode().iloc[0]))

# 2b. negative Unit_Price -> positive (data-entry errors, not real negative sales)
clean["Unit_Price"] = clean["Unit_Price"].abs()

# 2c. recalculate Total_Sales from the corrected price * quantity
clean["Total_Sales"] = (clean["Unit_Price"] * clean["Quantity"]).round(2)

# 2d. standardise category casing + date format
clean["Category"] = clean["Category"].str.title()
clean["Order_Date"] = pd.to_datetime(clean["Order_Date"]).dt.strftime("%Y-%m-%d")

clean_path = os.path.join(CLEAN_DIR, "retail_store_sales_cleaned.csv")
clean.to_csv(clean_path, index=False)

# ----------------------------------------------------------------------
# 3. KEY METRICS
# ----------------------------------------------------------------------
clean["Order_Date"] = pd.to_datetime(clean["Order_Date"])
total_sales = clean["Total_Sales"].sum()
n_orders = len(clean)
aov = total_sales / n_orders

monthly = (clean.groupby(clean["Order_Date"].dt.to_period("M"))
                .agg(sales=("Total_Sales", "sum"), orders=("Order_ID", "count")))
best_month = monthly["sales"].idxmax().strftime("%B")

cat_rev = clean.groupby("Category")["Total_Sales"].sum().sort_values(ascending=False)
cat_cnt = clean.groupby("Category")["Order_ID"].count()
reg_rev = clean.groupby("Region")["Total_Sales"].sum().sort_values(ascending=False)

print(f"\nKPIs -> Total: {total_sales:,.0f} | Orders: {n_orders} "
      f"| AOV: {aov:,.0f} | Best month: {best_month}")
print("Category revenue share %:", (cat_rev / total_sales * 100).round(1).to_dict())
print("Category order counts:", cat_cnt.to_dict())
print("Region sales:", reg_rev.round(0).to_dict())

# ----------------------------------------------------------------------
# 4. EXCEL DASHBOARD  (KPI cards + native charts)
# ----------------------------------------------------------------------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows

PURPLE, LIGHT, DARK, WHITE = "6B4E8E", "EDE7F4", "2D2333", "FFFFFF"
thin = Border(*[Side(style="thin", color="D9D2E3")] * 4)

wb = Workbook()

# data sheet
ws_data = wb.active
ws_data.title = "Data"
for r in dataframe_to_rows(clean, index=False, header=True):
    ws_data.append(r)
for c in ws_data[1]:
    c.font = Font(bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=PURPLE)

# summary sheets feeding the charts
ws_m = wb.create_sheet("Monthly Trend")
ws_m.append(["Month", "Sales", "Orders"])
for p, row in monthly.iterrows():
    ws_m.append([str(p), row["sales"], row["orders"]])
for c in ws_m[1]:
    c.font = Font(bold=True)

cat_df = (clean.groupby("Category")
             .agg(Revenue=("Total_Sales", "sum"), Orders=("Order_ID", "count"))
             .reset_index().sort_values("Revenue", ascending=False))
ws_c = wb.create_sheet("Category Summary")
ws_c.append(["Category", "Revenue", "Orders"])
for _, r in cat_df.iterrows():
    ws_c.append([r["Category"], r["Revenue"], r["Orders"]])
for c in ws_c[1]:
    c.font = Font(bold=True)

reg_df = (clean.groupby("Region")
             .agg(Revenue=("Total_Sales", "sum"), Orders=("Order_ID", "count"))
             .reset_index().sort_values("Revenue", ascending=False))
ws_r = wb.create_sheet("Region Summary")
ws_r.append(["Region", "Revenue", "Orders"])
for _, r in reg_df.iterrows():
    ws_r.append([r["Region"], r["Revenue"], r["Orders"]])
for c in ws_r[1]:
    c.font = Font(bold=True)

# dashboard sheet
ws = wb.create_sheet("Dashboard", 0)
ws.sheet_view.showGridLines = False
ws["B2"] = "Retail Store Sales Dashboard — 2025"
ws["B2"].font = Font(size=16, bold=True, color=DARK)
ws["B3"] = "One-year view · 4 categories · 4 regions · cleaned dataset"
ws["B3"].font = Font(size=10, italic=True, color="7A7480")

# KPI cards
kpis = [("Total Sales", f"{total_sales:,.0f}"), ("Total Orders", f"{n_orders:,}"),
        ("Avg Order Value", f"{aov:,.0f}"), ("Best Month", best_month)]
col = 2
for label, val in kpis:
    ws.cell(row=5, column=col, value=label)
    ws.cell(row=6, column=col, value=val)
    for rr in (5, 6):
        cell = ws.cell(row=rr, column=col)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.border = thin
        cell.alignment = Alignment(horizontal="center")
    ws.cell(row=5, column=col).font = Font(size=9, bold=True, color="7A7480")
    ws.cell(row=6, column=col).font = Font(size=15, bold=True, color=PURPLE)
    ws.column_dimensions[chr(64 + col)].width = 18
    col += 2

# charts
line = LineChart()
line.title = "Monthly Sales Trend (2025)"
line.height, line.width = 7, 15
line.add_data(Reference(ws_m, min_col=2, min_row=1, max_row=13), titles_from_data=True)
line.set_categories(Reference(ws_m, min_col=1, min_row=2, max_row=13))
line.y_axis.title = "Sales"
ws.add_chart(line, "B9")

bar = BarChart()
bar.type, bar.title = "col", "Revenue by Category"
bar.height, bar.width = 7, 9
bar.add_data(Reference(ws_c, min_col=2, min_row=1, max_row=5), titles_from_data=True)
bar.set_categories(Reference(ws_c, min_col=1, min_row=2, max_row=5))
bar.legend = None
ws.add_chart(bar, "J9")

bar2 = BarChart()
bar2.type, bar2.title = "col", "Revenue by Region"
bar2.height, bar2.width = 7, 8
bar2.add_data(Reference(ws_r, min_col=2, min_row=1, max_row=5), titles_from_data=True)
bar2.set_categories(Reference(ws_r, min_col=1, min_row=2, max_row=5))
bar2.legend = None
ws.add_chart(bar2, "B24")

# category x region matrix
ws["J24"] = "Category × Region (Revenue)"
ws["J24"].font = Font(bold=True, size=11, color=DARK)
piv = clean.pivot_table(index="Category", columns="Region",
                        values="Total_Sales", aggfunc="sum").round(0)
r0 = 25
ws.cell(row=r0, column=10, value="Category").font = Font(bold=True)
for j, reg in enumerate(piv.columns):
    ws.cell(row=r0, column=11 + j, value=reg).font = Font(bold=True)
for i, (cat, rowv) in enumerate(piv.iterrows()):
    ws.cell(row=r0 + 1 + i, column=10, value=cat)
    for j, v in enumerate(rowv.values):
        ws.cell(row=r0 + 1 + i, column=11 + j, value=float(v))

xlsx_path = os.path.join(DASH_DIR, "retail_sales_dashboard.xlsx")
wb.save(xlsx_path)
print("\nSaved dashboard:", xlsx_path)

# ----------------------------------------------------------------------
# 5. PNG PREVIEW OF THE DASHBOARD LAYOUT
# ----------------------------------------------------------------------
plt.rcParams.update({"figure.dpi": 110,
                     "axes.spines.top": False, "axes.spines.right": False})

fig = plt.figure(figsize=(13, 8))
gs = fig.add_gridspec(3, 2, height_ratios=[0.8, 2, 2], hspace=0.5, wspace=0.25)
fig.suptitle("Retail Store Sales Dashboard — 2025",
             fontsize=16, fontweight="bold", x=0.06, ha="left")

axk = fig.add_subplot(gs[0, :]); axk.axis("off")
for i, (label, val) in enumerate(kpis):
    axk.text(0.02 + i * 0.26, 0.5, f"{label}\n{val}",
             fontsize=12, fontweight="bold", color="#6B4E8E",
             bbox=dict(boxstyle="round,pad=0.6", fc="#EDE7F4", ec="#D9D2E3"),
             va="center")

ax1 = fig.add_subplot(gs[1, 0])
ax1.plot(range(12), monthly["sales"].values, marker="o", color="#6B4E8E")
ax1.set_xticks(range(12))
ax1.set_xticklabels([p.strftime("%b") for p in monthly.index], fontsize=8)
ax1.set_title("Monthly Sales Trend", fontsize=11)

ax2 = fig.add_subplot(gs[1, 1])
ax2.bar(cat_df["Category"], cat_df["Revenue"],
        color=["#6B4E8E", "#8F74AE", "#B39DCC", "#D8CBE8"])
ax2.set_title("Revenue by Category", fontsize=11)
ax2.tick_params(axis="x", rotation=15, labelsize=8)

ax3 = fig.add_subplot(gs[2, 0])
ax3.bar(reg_df["Region"], reg_df["Revenue"], color="#8F74AE")
ax3.set_title("Revenue by Region", fontsize=11)

ax4 = fig.add_subplot(gs[2, 1])
ax4.imshow(piv.values, cmap="Purples", aspect="auto")
ax4.set_xticks(range(4)); ax4.set_xticklabels(piv.columns, fontsize=8)
ax4.set_yticks(range(4)); ax4.set_yticklabels(piv.index, fontsize=8)
for i in range(4):
    for j in range(4):
        ax4.text(j, i, f"{piv.values[i, j]:,.0f}",
                 ha="center", va="center", fontsize=7)
ax4.set_title("Category × Region (Revenue)", fontsize=11)

png_path = os.path.join(DASH_DIR, "dashboard_preview.png")
plt.savefig(png_path, bbox_inches="tight")
plt.show()
print("Saved preview:", png_path)
